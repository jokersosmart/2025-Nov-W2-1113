"""
Unit tests for ExcelGenerator service
"""
from io import BytesIO
from datetime import datetime, timezone

import pytest
import openpyxl

from src.models.comment import Comment
from src.models.post import Platform, Post
from src.services.excel_generator import ExcelGenerator, ExcelGeneratorError


class TestExcelGenerator:
    """Test suite for ExcelGenerator service"""

    @pytest.fixture
    def generator(self) -> ExcelGenerator:
        """Create ExcelGenerator instance"""
        return ExcelGenerator()

    @pytest.fixture
    def sample_post(self) -> Post:
        """Create sample post for testing"""
        return Post(
            platform=Platform.FACEBOOK,
            post_url="https://www.facebook.com/test/posts/123456",
            post_time="2025-11-13T10:30:00Z",
            post_content="這是測試貼文內容",
            likes_count=150,
            comments_count=3,
        )

    @pytest.fixture
    def sample_comments(self) -> list[Comment]:
        """Create sample comments for testing"""
        return [
            Comment(
                comment_id="550e8400-e29b-41d4-a716-446655440001",
                post_url="https://www.facebook.com/test/posts/123456",
                comment_time="2025-11-13T11:00:00Z",
                commenter_id="王小明 (12345678)",
                comment_content="很棒的分享!",
                reply_window="張專員",
                reply_content="感謝您的支持!",
                customer_notes="已確認",
                generated_reply="",
            ),
            Comment(
                comment_id="550e8400-e29b-41d4-a716-446655440002",
                post_url="https://www.facebook.com/test/posts/123456",
                comment_time="2025-11-13T11:15:00Z",
                commenter_id="李小華 (87654321)",
                comment_content="請問有其他顏色嗎?",
                reply_window="陳客服",
                reply_content="目前有紅色與藍色可選",
                customer_notes="待回覆",
                generated_reply="",
            ),
            Comment(
                comment_id="550e8400-e29b-41d4-a716-446655440003",
                post_url="https://www.facebook.com/test/posts/123456",
                comment_time="2025-11-13T12:00:00Z",
                commenter_id="林美美 (11223344)",
                comment_content="已下單,期待收到商品!",
                reply_window="",
                reply_content="",
                customer_notes="",
                generated_reply="",
            ),
        ]

    def test_generator_initialization(self, generator: ExcelGenerator) -> None:
        """Test that ExcelGenerator can be initialized"""
        assert generator is not None
        assert generator.HEADERS == [
            "發文時間",
            "貼文連結",
            "貼文內容",
            "按讚數",
            "留言總數",
            "留言時間",
            "留言者 ID",
            "留言內容",
            "回覆窗口",
            "回覆內容",
            "客戶確認",
            "生成回覆",
        ]

    def test_generate_returns_bytesio(
        self, generator: ExcelGenerator, sample_post: Post, sample_comments: list[Comment]
    ) -> None:
        """Test that generate returns BytesIO object"""
        result = generator.generate(sample_post, sample_comments)
        assert isinstance(result, BytesIO)
        assert result.tell() == 0  # Should be seeked to start

    def test_generate_creates_valid_excel(
        self, generator: ExcelGenerator, sample_post: Post, sample_comments: list[Comment]
    ) -> None:
        """Test that generated file is a valid Excel workbook"""
        excel_file = generator.generate(sample_post, sample_comments)

        # Try to load with openpyxl
        workbook = openpyxl.load_workbook(excel_file)
        assert workbook is not None
        assert workbook.active is not None

    def test_excel_has_correct_headers(
        self, generator: ExcelGenerator, sample_post: Post, sample_comments: list[Comment]
    ) -> None:
        """Test that Excel file has correct header row"""
        excel_file = generator.generate(sample_post, sample_comments)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active

        # Check header row
        expected_headers = generator.HEADERS
        actual_headers = [cell.value for cell in sheet[1]]

        assert actual_headers == expected_headers

    def test_excel_has_correct_data_rows(
        self, generator: ExcelGenerator, sample_post: Post, sample_comments: list[Comment]
    ) -> None:
        """Test that Excel file has correct data rows"""
        excel_file = generator.generate(sample_post, sample_comments)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active

        # Should have header + 3 data rows
        assert sheet.max_row == 4

        # Check first data row (row 2)
        # Note: datetime values are stored as strings after conversion
        assert str(sheet["A2"].value).startswith("2025-11-13")  # 發文時間
        assert (
            sheet["B2"].value == "https://www.facebook.com/test/posts/123456"
        )  # 貼文連結
        assert sheet["C2"].value == "這是測試貼文內容"  # 貼文內容
        assert sheet["D2"].value == 150  # 按讚數
        assert sheet["E2"].value == 3  # 留言總數
        assert str(sheet["F2"].value).startswith("2025-11-13")  # 留言時間
        assert sheet["G2"].value == "王小明 (12345678)"  # 留言者 ID
        assert sheet["H2"].value == "很棒的分享!"  # 留言內容
        assert sheet["I2"].value == "張專員"  # 回覆窗口
        assert sheet["J2"].value == "感謝您的支持!"  # 回覆內容
        assert sheet["K2"].value == "已確認"  # 客戶確認
        assert sheet["L2"].value in ["", None]  # 生成回覆 (empty string becomes None in Excel)

        # Check third data row (row 4) - with empty fields
        assert sheet["I4"].value in ["", None]  # 回覆窗口 (empty)
        assert sheet["J4"].value in ["", None]  # 回覆內容 (empty)
        assert sheet["K4"].value in ["", None]  # 客戶確認 (empty)

    def test_excel_handles_empty_comments(
        self, generator: ExcelGenerator, sample_post: Post
    ) -> None:
        """Test that Excel generation works with empty comments list"""
        excel_file = generator.generate(sample_post, [])
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active

        # Should only have header row
        assert sheet.max_row == 1

        # Check header is still present
        assert [cell.value for cell in sheet[1]] == generator.HEADERS

    def test_excel_post_data_repeats_for_each_comment(
        self, generator: ExcelGenerator, sample_post: Post, sample_comments: list[Comment]
    ) -> None:
        """Test that post data is repeated for each comment row"""
        excel_file = generator.generate(sample_post, sample_comments)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active

        # Check that post data is same in all rows
        for row_idx in range(2, 5):  # Rows 2, 3, 4
            assert str(sheet[f"A{row_idx}"].value).startswith("2025-11-13")  # 發文時間
            assert (
                sheet[f"B{row_idx}"].value
                == "https://www.facebook.com/test/posts/123456"
            )  # 貼文連結
            assert sheet[f"C{row_idx}"].value == "這是測試貼文內容"  # 貼文內容
            assert sheet[f"D{row_idx}"].value == 150  # 按讚數
            assert sheet[f"E{row_idx}"].value == 3  # 留言總數

    def test_generate_filename_facebook(self) -> None:
        """Test filename generation for Facebook"""
        filename = ExcelGenerator.generate_filename("facebook")
        assert filename.startswith("facebook_comments_")
        assert filename.endswith(".xlsx")
        # Check timestamp format YYYYMMDD_HHMMSS
        parts = filename.replace("facebook_comments_", "").replace(".xlsx", "")
        assert len(parts) == 15  # YYYYMMDD_HHMMSS
        assert "_" in parts

    def test_generate_filename_instagram(self) -> None:
        """Test filename generation for Instagram"""
        filename = ExcelGenerator.generate_filename("instagram")
        assert filename.startswith("instagram_comments_")
        assert filename.endswith(".xlsx")

    def test_generate_filename_case_insensitive(self) -> None:
        """Test filename generation handles uppercase platform names"""
        filename = ExcelGenerator.generate_filename("FACEBOOK")
        assert filename.startswith("facebook_comments_")  # Should be lowercase

    def test_generate_filename_unique_timestamps(self) -> None:
        """Test that consecutive filename generations have different timestamps"""
        import time

        filename1 = ExcelGenerator.generate_filename("facebook")
        time.sleep(1.1)  # Wait more than 1 second
        filename2 = ExcelGenerator.generate_filename("facebook")

        assert filename1 != filename2

    def test_excel_headers_have_bold_font(
        self, generator: ExcelGenerator, sample_post: Post, sample_comments: list[Comment]
    ) -> None:
        """Test that header cells have bold font"""
        excel_file = generator.generate(sample_post, sample_comments)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active

        # Check all header cells are bold
        for cell in sheet[1]:
            assert cell.font.bold is True

    def test_excel_column_widths_adjusted(
        self, generator: ExcelGenerator, sample_post: Post, sample_comments: list[Comment]
    ) -> None:
        """Test that column widths are adjusted"""
        excel_file = generator.generate(sample_post, sample_comments)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active

        # Check that some columns have custom widths
        # Column B (貼文連結) should be wider
        assert sheet.column_dimensions["B"].width == 50
        # Column A (發文時間) should be moderate
        assert sheet.column_dimensions["A"].width == 20
        # Column D (按讚數) should be narrow
        assert sheet.column_dimensions["D"].width == 10

"""
Contract tests for POST /api/export endpoint

These tests define the expected behavior of the export API endpoint
before implementation (TDD approach).
"""
import pytest
from fastapi.testclient import TestClient
from io import BytesIO
import openpyxl

from main import app


class TestExportContract:
    """Contract tests for export endpoint"""

    @pytest.fixture
    def client(self) -> TestClient:
        """Create test client"""
        return TestClient(app)

    @pytest.fixture
    def sample_export_data(self) -> dict:
        """Sample data for export testing"""
        return {
            "post": {
                "platform": "facebook",
                "post_url": "https://www.facebook.com/test/posts/123456",
                "post_time": "2025-11-13T10:30:00Z",
                "post_content": "這是測試貼文內容",
                "likes_count": 150,
                "comments_count": 3
            },
            "comments": [
                {
                    "comment_id": "550e8400-e29b-41d4-a716-446655440001",
                    "post_url": "https://www.facebook.com/test/posts/123456",
                    "comment_time": "2025-11-13T11:00:00Z",
                    "commenter_id": "王小明 (12345678)",
                    "comment_content": "很棒的分享!",
                    "reply_window": "張專員",
                    "reply_content": "感謝您的支持!",
                    "customer_notes": "已確認",
                    "generated_reply": ""
                },
                {
                    "comment_id": "550e8400-e29b-41d4-a716-446655440002",
                    "post_url": "https://www.facebook.com/test/posts/123456",
                    "comment_time": "2025-11-13T11:15:00Z",
                    "commenter_id": "李小華 (87654321)",
                    "comment_content": "請問有其他顏色嗎?",
                    "reply_window": "陳客服",
                    "reply_content": "目前有紅色與藍色可選",
                    "customer_notes": "待回覆",
                    "generated_reply": ""
                },
                {
                    "comment_id": "550e8400-e29b-41d4-a716-446655440003",
                    "post_url": "https://www.facebook.com/test/posts/123456",
                    "comment_time": "2025-11-13T12:00:00Z",
                    "commenter_id": "林美美 (11223344)",
                    "comment_content": "已下單,期待收到商品!",
                    "reply_window": "",
                    "reply_content": "",
                    "customer_notes": "",
                    "generated_reply": ""
                }
            ]
        }

    def test_export_endpoint_exists(self, client: TestClient, sample_export_data: dict) -> None:
        """Test that POST /api/export endpoint is accessible"""
        response = client.post("/api/export", json=sample_export_data)
        # Should not return 404
        assert response.status_code != 404

    def test_export_requires_post_data(self, client: TestClient) -> None:
        """Test that post data is required in request body"""
        response = client.post("/api/export", json={"comments": []})
        assert response.status_code in [400, 422]  # Accept both for Pydantic validation
        data = response.json()
        assert "post" in str(data).lower() or "detail" in data

    def test_export_requires_comments_data(self, client: TestClient) -> None:
        """Test that comments data is required in request body"""
        response = client.post(
            "/api/export",
            json={
                "post": {
                    "platform": "facebook",
                    "post_url": "https://www.facebook.com/test/posts/123",
                    "post_time": "2025-11-13T10:00:00Z",
                    "post_content": "測試",
                    "likes_count": 0,
                    "comments_count": 0
                }
            }
        )
        assert response.status_code in [400, 422]  # Accept both for Pydantic validation
        data = response.json()
        assert "comments" in str(data).lower() or "detail" in data

    def test_export_validates_platform(self, client: TestClient, sample_export_data: dict) -> None:
        """Test that invalid platform is rejected"""
        invalid_data = sample_export_data.copy()
        invalid_data["post"]["platform"] = "twitter"
        response = client.post("/api/export", json=invalid_data)
        assert response.status_code in [400, 422]  # Accept both for Pydantic validation
        data = response.json()
        assert "platform" in str(data).lower() or "detail" in data

    def test_export_returns_excel_file(self, client: TestClient, sample_export_data: dict) -> None:
        """Test that response is an Excel file"""
        response = client.post("/api/export", json=sample_export_data)
        if response.status_code != 200:
            print(f"Response status: {response.status_code}")
            print(f"Response body: {response.json()}")
        assert response.status_code == 200
        # Check content type
        assert response.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

    def test_export_returns_correct_filename(self, client: TestClient, sample_export_data: dict) -> None:
        """Test that response includes correct filename with platform and timestamp"""
        response = client.post("/api/export", json=sample_export_data)
        assert response.status_code == 200
        # Check Content-Disposition header
        content_disposition = response.headers.get("content-disposition", "")
        assert "attachment" in content_disposition
        assert "facebook_comments_" in content_disposition
        assert ".xlsx" in content_disposition

    def test_export_file_is_valid_excel(self, client: TestClient, sample_export_data: dict) -> None:
        """Test that exported file is a valid Excel workbook"""
        response = client.post("/api/export", json=sample_export_data)
        assert response.status_code == 200
        
        # Try to load the file with openpyxl
        excel_file = BytesIO(response.content)
        try:
            workbook = openpyxl.load_workbook(excel_file)
            assert workbook is not None
        except Exception as e:
            pytest.fail(f"Failed to load Excel file: {e}")

    def test_export_excel_has_correct_columns(self, client: TestClient, sample_export_data: dict) -> None:
        """Test that Excel file contains all required columns in Chinese"""
        response = client.post("/api/export", json=sample_export_data)
        assert response.status_code == 200
        
        excel_file = BytesIO(response.content)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active
        
        # Check header row (row 1)
        expected_headers = [
            "發文時間",
            "貼文連結",
            "貼文內容",
            "按讚數",
            "留言總數",
            "留言時間",
            "留言者 ID",
            "留言內容",
            "回覆窗口",
            "回覆內容",
            "客戶確認",
            "生成回覆"
        ]
        
        actual_headers = [cell.value for cell in sheet[1]]
        assert actual_headers == expected_headers

    def test_export_excel_contains_correct_data(self, client: TestClient, sample_export_data: dict) -> None:
        """Test that Excel file contains correct data rows"""
        response = client.post("/api/export", json=sample_export_data)
        assert response.status_code == 200
        
        excel_file = BytesIO(response.content)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active
        
        # Should have header + 3 data rows
        assert sheet.max_row == 4
        
        # Check first data row (row 2)
        assert str(sheet["A2"].value).startswith("2025-11-13")  # 發文時間
        assert sheet["B2"].value == "https://www.facebook.com/test/posts/123456"  # 貼文連結
        assert sheet["C2"].value == "這是測試貼文內容"  # 貼文內容
        assert sheet["D2"].value == 150  # 按讚數
        assert sheet["E2"].value == 3  # 留言總數
        assert str(sheet["F2"].value).startswith("2025-11-13")  # 留言時間
        assert sheet["G2"].value == "王小明 (12345678)"  # 留言者 ID
        assert sheet["H2"].value == "很棒的分享!"  # 留言內容
        assert sheet["I2"].value == "張專員"  # 回覆窗口
        assert sheet["J2"].value == "感謝您的支持!"  # 回覆內容
        assert sheet["K2"].value == "已確認"  # 客戶確認
        assert sheet["L2"].value in ["", None]  # 生成回覆

    def test_export_handles_empty_comments(self, client: TestClient) -> None:
        """Test that export works with empty comments list"""
        data = {
            "post": {
                "platform": "instagram",
                "post_url": "https://www.instagram.com/p/ABC123/",
                "post_time": "2025-11-13T10:00:00Z",
                "post_content": "測試貼文",
                "likes_count": 50,
                "comments_count": 0
            },
            "comments": []
        }
        response = client.post("/api/export", json=data)
        assert response.status_code == 200
        
        excel_file = BytesIO(response.content)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active
        
        # Should only have header row
        assert sheet.max_row == 1

    def test_export_handles_partial_comment_data(self, client: TestClient) -> None:
        """Test that export handles comments with missing optional fields"""
        data = {
            "post": {
                "platform": "facebook",
                "post_url": "https://www.facebook.com/test/posts/999",
                "post_time": "2025-11-13T10:00:00Z",
                "post_content": "測試",
                "likes_count": 10,
                "comments_count": 1
            },
            "comments": [
                {
                    "comment_id": "550e8400-e29b-41d4-a716-446655440099",
                    "post_url": "https://www.facebook.com/test/posts/999",
                    "comment_time": "2025-11-13T11:00:00Z",
                    "commenter_id": "測試者 (99999999)",
                    "comment_content": "測試留言",
                    "reply_window": "",  # Empty fields should be handled
                    "reply_content": "",
                    "customer_notes": "",
                    "generated_reply": ""
                }
            ]
        }
        response = client.post("/api/export", json=data)
        assert response.status_code == 200
        
        excel_file = BytesIO(response.content)
        workbook = openpyxl.load_workbook(excel_file)
        sheet = workbook.active
        
        # Check that empty fields are present
        assert sheet["I2"].value in ["", None]  # 回覆窗口
        assert sheet["J2"].value in ["", None]  # 回覆內容
        assert sheet["K2"].value in ["", None]  # 客戶確認

    def test_export_instagram_filename(self, client: TestClient) -> None:
        """Test that Instagram exports have correct filename prefix"""
        data = {
            "post": {
                "platform": "instagram",
                "post_url": "https://www.instagram.com/p/ABC123/",
                "post_time": "2025-11-13T10:00:00Z",
                "post_content": "測試",
                "likes_count": 100,
                "comments_count": 0
            },
            "comments": []
        }
        response = client.post("/api/export", json=data)
        assert response.status_code == 200
        
        content_disposition = response.headers.get("content-disposition", "")
        assert "instagram_comments_" in content_disposition
        assert ".xlsx" in content_disposition

    def test_export_content_type(self, client: TestClient, sample_export_data: dict) -> None:
        """Test that endpoint requires JSON content type"""
        response = client.post(
            "/api/export",
            data="invalid data",
            headers={"Content-Type": "text/plain"},
        )
        # Should reject non-JSON content
        assert response.status_code in [400, 415, 422]

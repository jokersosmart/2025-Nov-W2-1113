"""
Excel 匯出服務

使用 openpyxl 生成 .xlsx 檔案,包含貼文與留言資料
"""
from datetime import datetime
from io import BytesIO
from typing import List

import openpyxl
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.worksheet import Worksheet

from src.models.comment import Comment
from src.models.post import Post


class ExcelGeneratorError(Exception):
    """Excel 生成服務錯誤"""

    pass


class ExcelGenerator:
    """
    Excel 生成服務

    將貼文與留言資料生成為 .xlsx 格式檔案
    """

    # 欄位標題 (中文)
    HEADERS = [
        "發文時間",
        "貼文連結",
        "貼文內容",
        "按讚數",
        "留言總數",
        "留言時間",
        "留言者 ID",
        "留言內容",
        "回覆窗口",
        "回覆內容",
        "客戶確認",
        "生成回覆",
    ]

    def __init__(self) -> None:
        """初始化 Excel 生成器"""
        pass

    def generate(self, post: Post, comments: List[Comment]) -> BytesIO:
        """
        生成 Excel 檔案

        Args:
            post: 貼文資料
            comments: 留言列表

        Returns:
            BytesIO: Excel 檔案二進位資料

        Raises:
            ExcelGeneratorError: 生成過程發生錯誤
        """
        try:
            # 建立工作簿與工作表
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            if sheet is None:
                raise ExcelGeneratorError("無法建立工作表")

            # 設定標題列
            self._write_headers(sheet)

            # 寫入資料列
            self._write_data_rows(sheet, post, comments)

            # 調整欄寬
            self._adjust_column_widths(sheet)

            # 儲存至 BytesIO
            output = BytesIO()
            workbook.save(output)
            output.seek(0)

            return output

        except Exception as e:
            raise ExcelGeneratorError(f"Excel 生成失敗: {str(e)}") from e

    def _write_headers(self, sheet: Worksheet) -> None:
        """
        寫入標題列

        Args:
            sheet: 工作表
        """
        for col_idx, header in enumerate(self.HEADERS, start=1):
            cell = sheet.cell(row=1, column=col_idx, value=header)
            # 設定粗體字型
            cell.font = Font(bold=True)
            # 設定對齊方式
            cell.alignment = Alignment(horizontal="center", vertical="center")

    def _write_data_rows(
        self, sheet: Worksheet, post: Post, comments: List[Comment]
    ) -> None:
        """
        寫入資料列

        Args:
            sheet: 工作表
            post: 貼文資料
            comments: 留言列表
        """
        # 如果沒有留言,只寫入標題列
        if not comments:
            return

        # 每一則留言寫入一列
        for row_idx, comment in enumerate(comments, start=2):
            # 貼文欄位 (每列都重複)
            sheet.cell(row=row_idx, column=1, value=str(post.post_time))  # 發文時間
            sheet.cell(row=row_idx, column=2, value=str(post.post_url))  # 貼文連結
            sheet.cell(row=row_idx, column=3, value=post.post_content)  # 貼文內容
            sheet.cell(row=row_idx, column=4, value=post.likes_count)  # 按讚數
            sheet.cell(row=row_idx, column=5, value=post.comments_count)  # 留言總數

            # 留言欄位
            sheet.cell(row=row_idx, column=6, value=str(comment.comment_time))  # 留言時間
            sheet.cell(row=row_idx, column=7, value=comment.commenter_id)  # 留言者 ID
            sheet.cell(row=row_idx, column=8, value=comment.comment_content)  # 留言內容
            sheet.cell(row=row_idx, column=9, value=comment.reply_window or "")  # 回覆窗口
            sheet.cell(
                row=row_idx, column=10, value=comment.reply_content or ""
            )  # 回覆內容
            sheet.cell(
                row=row_idx, column=11, value=comment.customer_notes or ""
            )  # 客戶確認
            sheet.cell(
                row=row_idx, column=12, value=comment.generated_reply or ""
            )  # 生成回覆

    def _adjust_column_widths(self, sheet: Worksheet) -> None:
        """
        調整欄寬以符合內容

        Args:
            sheet: 工作表
        """
        # 預設欄寬設定 (根據欄位內容類型)
        column_widths = {
            1: 20,  # 發文時間
            2: 50,  # 貼文連結
            3: 40,  # 貼文內容
            4: 10,  # 按讚數
            5: 10,  # 留言總數
            6: 20,  # 留言時間
            7: 30,  # 留言者 ID
            8: 40,  # 留言內容
            9: 15,  # 回覆窗口
            10: 40,  # 回覆內容
            11: 15,  # 客戶確認
            12: 40,  # 生成回覆
        }

        for col_idx, width in column_widths.items():
            sheet.column_dimensions[openpyxl.utils.get_column_letter(col_idx)].width = (
                width
            )

    @staticmethod
    def generate_filename(platform: str) -> str:
        """
        生成檔名

        Args:
            platform: 平台名稱 (facebook/instagram)

        Returns:
            str: 檔名,格式為 {platform}_comments_YYYYMMDD_HHMMSS.xlsx

        Example:
            facebook_comments_20251113_143052.xlsx
        """
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        return f"{platform.lower()}_comments_{timestamp}.xlsx"
